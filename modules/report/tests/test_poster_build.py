# -*- coding: utf-8 -*-
"""Synthetic-form tests for the report module's poster CLI."""
from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

from module_registry import ModuleRegistry


_HERE = Path(__file__).resolve().parent
_MODULE_ROOT = _HERE.parent
_REPO_ROOT = _MODULE_ROOT.parents[1]
_SCRIPTS = _MODULE_ROOT / "scripts"
POSTER_BUILD = _SCRIPTS / "poster_build.py"
POSTER_VERIFY = _SCRIPTS / "poster_verify.py"

EMU_PER_CM = 360000
BOX_ALPHA = "BoxAlpha"
BOX_BETA = "BoxBeta"


def _poster_deps_available() -> bool:
    try:
        import pptx  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError:
        return False
    return True


def _run(script: Path, args: list[str], env=None) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(script), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if str(getattr(layout, "name", "")).lower() == "blank":
            return layout
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


def _write_form(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(_blank_layout(prs))

    def add_box(name, left_cm, top_cm, width_cm, height_cm):
        shape = slide.shapes.add_textbox(
            Emu(int(left_cm * EMU_PER_CM)),
            Emu(int(top_cm * EMU_PER_CM)),
            Emu(int(width_cm * EMU_PER_CM)),
            Emu(int(height_cm * EMU_PER_CM)),
        )
        shape.name = name
        return shape

    add_box(BOX_ALPHA, 0.5, 0.4, 12.0, 16.0)
    add_box(BOX_BETA, 13.5, 0.4, 12.0, 16.0)
    prs.save(path)


def _write_png(path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (16, 16), (40, 80, 160)).save(path)


def _content_text(*, alpha: str, beta: str, with_fig: bool) -> str:
    fig = ""
    if with_fig:
        fig = '\n\n[[FIG file="plot.png" caption="Synthetic plot"]]\n'
    return (
        "# TITLE: Synthetic Poster\n"
        "# AUTHORS: Demo Authors\n"
        "box_map:\n"
        f'{{"alpha": ["{BOX_ALPHA}", null], "beta": ["{BOX_BETA}", null]}}\n'
        "\n"
        "## BOX: alpha\n"
        f"{alpha}{fig}\n"
        "## BOX: beta\n"
        f"{beta}\n"
    )


def _registry(tmp_path: Path) -> ModuleRegistry:
    enabled = tmp_path / "enabled.yaml"
    enabled.write_text(
        "schema: rigorloom-enabled-modules/v1\n"
        "enabled: [report, style]\n",
        encoding="utf-8",
    )
    return ModuleRegistry(
        _REPO_ROOT / "modules",
        enabled_file=enabled,
        pyproject=_REPO_ROOT / "pyproject.toml",
    )


@pytest.mark.skipif(not _poster_deps_available(),
                    reason="python-pptx and Pillow are not installed")
def test_build_verify_then_overflow(tmp_path: Path) -> None:
    from pptx import Presentation

    form = tmp_path / "form.pptx"
    out = tmp_path / "poster.pptx"
    figures = tmp_path / "figures"
    figures.mkdir()
    _write_png(figures / "plot.png")
    _write_form(form)
    form_hash = _sha256(form)
    form_mtime = form.stat().st_mtime

    content = tmp_path / "content.md"
    content.write_text(
        _content_text(
            alpha="Alpha body sentence for the first box.",
            beta="Beta body sentence for the second box.",
            with_fig=True,
        ),
        encoding="utf-8",
    )

    built = _run(POSTER_BUILD, [
        "--content", str(content),
        "--figures", str(figures),
        "--form", str(form),
        "--out", str(out),
        "--min-pt", "24",
    ])
    assert built.returncode == 0, built.stdout + built.stderr
    payload = json.loads(built.stdout)
    assert payload["ok"] is True
    assert out.is_file()
    assert _sha256(form) == form_hash

    prs = Presentation(out)
    by_name = {sh.name: sh for sh in prs.slides[0].shapes}
    assert "Alpha body sentence" in by_name[BOX_ALPHA].text_frame.text
    assert "Beta body sentence" in by_name[BOX_BETA].text_frame.text

    verified = _run(POSTER_VERIFY, [
        "--out", str(out),
        "--form", str(form),
        "--form-mtime", str(form_mtime),
        "--no-numbers",
        "--min-pictures", "1",
        "--min-pt", "24",
    ])
    assert verified.returncode == 0, verified.stdout + verified.stderr
    assert json.loads(verified.stdout)["ok"] is True

    overflow_content = tmp_path / "overflow.md"
    overflow_content.write_text(
        _content_text(
            alpha="overflow " * 2000,
            beta="Beta still short.",
            with_fig=False,
        ),
        encoding="utf-8",
    )
    overflow_out = tmp_path / "overflow.pptx"
    overflowed = _run(POSTER_BUILD, [
        "--content", str(overflow_content),
        "--figures", str(figures),
        "--form", str(form),
        "--out", str(overflow_out),
    ])
    assert overflowed.returncode != 0, overflowed.stdout
    overflow_json = json.loads(overflowed.stdout)
    assert overflow_json["ok"] is False
    dumped = json.dumps(overflow_json)
    assert "overflow" in dumped
    assert not overflow_out.exists()
    assert _sha256(form) == form_hash


def test_missing_dependency_exits_3(tmp_path: Path) -> None:
    blocker = tmp_path / "blocked"
    blocker.mkdir()
    (blocker / "pptx.py").write_text(
        "raise ImportError('blocked for poster_dependency_missing')\n",
        encoding="utf-8",
    )
    env = os.environ.copy()
    previous = env.get("PYTHONPATH", "")
    env["PYTHONPATH"] = str(blocker) + (os.pathsep + previous if previous else "")

    dummy = tmp_path / "dummy.md"
    dummy.write_text("# TITLE: t\n# AUTHORS: a\n", encoding="utf-8")
    form = tmp_path / "form.pptx"
    form.write_bytes(b"not-a-real-pptx")
    figs = tmp_path / "figs"
    figs.mkdir()

    built = _run(POSTER_BUILD, [
        "--content", str(dummy),
        "--figures", str(figs),
        "--form", str(form),
        "--out", str(tmp_path / "out.pptx"),
    ], env=env)
    assert built.returncode == 3, built.stdout + built.stderr
    payload = json.loads(built.stdout)
    assert payload["ok"] is False
    assert payload["error"]["code"] == "poster_dependency_missing"
    assert payload["error"]["extra"] == "poster"

    verified = _run(POSTER_VERIFY, [
        "--out", str(tmp_path / "out.pptx"),
        "--form", str(form),
        "--form-mtime", "0",
        "--no-numbers",
    ], env=env)
    assert verified.returncode == 3, verified.stdout + verified.stderr
    v_payload = json.loads(verified.stdout)
    assert v_payload["error"]["code"] == "poster_dependency_missing"


def test_registry_declares_poster_cli(tmp_path: Path) -> None:
    registry = _registry(tmp_path)
    commands = {row["command"]: row for row in registry.enabled_cli()}
    assert "poster" in commands
    assert "poster-verify" in commands
    assert commands["poster"]["module"] == "report"
    assert commands["poster-verify"]["module"] == "report"
    assert Path(commands["poster"]["script"]).resolve() == POSTER_BUILD.resolve()
    assert Path(commands["poster-verify"]["script"]).resolve() == POSTER_VERIFY.resolve()
