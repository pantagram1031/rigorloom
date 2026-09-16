#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill a poster PPTX form from a markdown content file.

Usage:
  python poster_build.py --content poster.md --figures <dir>
                         --form <form.pptx> --out <out.pptx>
                         [--box-map map.json] [--min-pt 24]

The form is copied, never modified. Body text is written at --min-pt
(default 24) in 맑은 고딕; figures sit below the text in each box.
Overflow is estimated from character metrics and is a hard failure —
fonts are never shrunk below --min-pt to force a fit.

Requires the optional extra: pip install ".[poster]"
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
from pathlib import Path
from types import SimpleNamespace


EMU_PER_CM = 360000
LINE_SPACING = 1.15
CHAR_W_CM_PER_PT = 0.03527
LINE_H_CM_PER_PT = 0.03527 * 1.3
FIG_GAP_CM = 0.6
SIDE_MARGIN_CM = 0.4

# Legacy six-box fallback so older content files still build. New forms
# must pass --box-map or a box_map: JSON header; this is not the only map.
DEFAULT_BOX_MAP = {
    "연구동기": ("한쪽 모서리가 잘린 사각형 5", "직사각형 11"),
    "이론적배경": ("한쪽 모서리가 잘린 사각형 35", "직사각형 26"),
    "연구과정A": ("한쪽 모서리가 잘린 사각형 3", "직사각형 30"),
    "연구과정B": ("한쪽 모서리가 잘린 사각형 39", None),
    "결론": ("한쪽 모서리가 잘린 사각형 40", "직사각형 27"),
    "참고문헌": ("한쪽 모서리가 잘린 사각형 41", "직사각형 29"),
}
DEFAULT_HEADER_SHAPE = "모서리가 둥근 직사각형 4"
DEFAULT_TITLE_PLACEHOLDER = "주제"
DEFAULT_AUTHORS_PLACEHOLDER = "저자"
DEFAULT_FONT = "맑은 고딕"
GUIDE_PATTERNS = [r"본문 내용은", r"맑은 고딕\s*\d+\s*pt"]

FIG_RE = re.compile(r'\[\[FIG\s+file="([^"]+)"\s+caption="([^"]+)"\s*\]\]')
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


class PosterError(Exception):
    def __init__(self, code: str, message: str, *, extra=None, failures=None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.extra = extra
        self.failures = failures


class JsonArgumentParser(argparse.ArgumentParser):
    def error(self, message):
        _die({"ok": False, "error": {"code": "usage", "message": message}}, 2)


def _print_json(obj: dict) -> None:
    data = json.dumps(obj, ensure_ascii=False) + "\n"
    try:
        sys.stdout.write(data)
    except UnicodeEncodeError:
        sys.stdout.buffer.write(data.encode("utf-8"))
    sys.stdout.flush()


def _die(payload: dict, code: int) -> None:
    _print_json(payload)
    raise SystemExit(code)


def _utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError, OSError):
            pass


def _missing_dep(exc: BaseException) -> None:
    _die(
        {
            "ok": False,
            "error": {
                "code": "poster_dependency_missing",
                "message": (
                    "python-pptx and Pillow are required "
                    f"({type(exc).__name__}: {exc}). "
                    'Install with pip install ".[poster]".'
                ),
                "extra": "poster",
            },
        },
        3,
    )


def import_poster_deps():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Pt
        from PIL import Image
    except ImportError as exc:
        _missing_dep(exc)
    return SimpleNamespace(
        Presentation=Presentation,
        RGBColor=RGBColor,
        MSO_ANCHOR=MSO_ANCHOR,
        PP_ALIGN=PP_ALIGN,
        Emu=Emu,
        Pt=Pt,
        Image=Image,
    )


def cm(deps, value: float):
    return deps.Emu(int(value * EMU_PER_CM))


def normalize_box_map(obj) -> dict:
    if not isinstance(obj, dict) or not obj:
        raise PosterError("invalid_box_map", "box_map must be a non-empty JSON object")
    out = {}
    for key, val in obj.items():
        if not isinstance(val, (list, tuple)) or not (1 <= len(val) <= 2):
            raise PosterError(
                "invalid_box_map",
                f"box_map[{key!r}] must be [content_shape, label_shape_or_null]",
            )
        box_name = val[0]
        label = val[1] if len(val) == 2 else None
        if not isinstance(box_name, str) or not box_name.strip():
            raise PosterError(
                "invalid_box_map",
                f"box_map[{key!r}] content shape must be a non-empty string",
            )
        if label is not None and (not isinstance(label, str) or not label.strip()):
            raise PosterError(
                "invalid_box_map",
                f"box_map[{key!r}] label must be a string or null",
            )
        out[str(key)] = (box_name, label)
    return out


def load_box_map_arg(value: str) -> dict:
    raw = value.strip()
    if raw.startswith("{"):
        try:
            payload = json.loads(raw)
        except json.JSONDecodeError as exc:
            raise PosterError("invalid_box_map", f"--box-map JSON is invalid: {exc}") from exc
        return normalize_box_map(payload)
    path = Path(value)
    try:
        text = path.read_text(encoding="utf-8")
    except OSError as exc:
        raise PosterError("invalid_box_map", f"--box-map file unreadable: {exc}") from exc
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        raise PosterError("invalid_box_map", f"--box-map file is not JSON: {exc}") from exc
    return normalize_box_map(payload)


def parse_box_map_header(header: str):
    match = re.search(r"^box_map:\s*", header, re.M)
    if not match:
        return None
    rest = header[match.end():].lstrip()
    if not rest.startswith("{"):
        raise PosterError(
            "invalid_box_map",
            "box_map: header must be a JSON object (see references/poster_content.md)",
        )
    try:
        payload, _end = json.JSONDecoder().raw_decode(rest)
    except json.JSONDecodeError as exc:
        raise PosterError("invalid_box_map", f"box_map: JSON is invalid: {exc}") from exc
    return normalize_box_map(payload)


def parse_content(path: str) -> dict:
    raw = Path(path).read_text(encoding="utf-8")
    title_m = re.search(r"^# TITLE:\s*(.+)$", raw, re.M)
    authors_m = re.search(r"^# AUTHORS:\s*(.+)$", raw, re.M)
    if not title_m or not authors_m:
        raise PosterError(
            "invalid_content",
            "content file must start with '# TITLE:' and '# AUTHORS:' lines",
        )
    header_m = re.search(r"^# HEADER_SHAPE:\s*(.+)$", raw, re.M)
    title_ph = re.search(r"^# TITLE_PLACEHOLDER:\s*(.+)$", raw, re.M)
    authors_ph = re.search(r"^# AUTHORS_PLACEHOLDER:\s*(.+)$", raw, re.M)
    parts = re.split(r"^## BOX:\s*(\S+)\s*$", raw, flags=re.M)
    header = parts[0] if parts else raw
    boxes = {}
    for i in range(1, len(parts) - 1, 2):
        key, body = parts[i], parts[i + 1]
        figs = [{"file": m.group(1), "caption": m.group(2)}
                for m in FIG_RE.finditer(body)]
        body = FIG_RE.sub("", body)
        paras = [p.strip() for p in body.split("\n\n") if p.strip()]
        boxes[key] = {"paras": paras, "figs": figs}
    return {
        "title": title_m.group(1).strip(),
        "authors": authors_m.group(1).strip(),
        "header_shape": header_m.group(1).strip() if header_m else None,
        "title_placeholder": title_ph.group(1).strip() if title_ph else None,
        "authors_placeholder": authors_ph.group(1).strip() if authors_ph else None,
        "box_map": parse_box_map_header(header),
        "boxes": boxes,
    }


def set_run_font(run, pt, deps, bold=False, font_name=DEFAULT_FONT):
    font = run.font
    font.name = font_name
    font.size = deps.Pt(pt)
    font.bold = bold
    font.color.rgb = deps.RGBColor(0, 0, 0)
    r_pr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = r_pr.find("{%s}%s" % (NS_A, tag))
        if el is None:
            el = r_pr.makeelement("{%s}%s" % (NS_A, tag), {})
            r_pr.append(el)
        el.set("typeface", font_name)


def fill_text(tf, paras, pt, deps, font_name=DEFAULT_FONT):
    tf.word_wrap = True
    tf.vertical_anchor = deps.MSO_ANCHOR.TOP
    tf.clear()
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = LINE_SPACING
        run = p.add_run()
        run.text = para
        set_run_font(run, pt, deps, font_name=font_name)


def eff_len(text: str) -> float:
    """Text length in full-width units: Hangul/CJK = 1, ASCII ~ 0.55."""
    return sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in text)


def est_text_h_cm(paras, width_cm, pt) -> float:
    cpl = max(1.0, width_cm / (pt * CHAR_W_CM_PER_PT))
    lines = sum(max(1, -(-int(eff_len(p) * 100) // int(cpl * 100)))
                for p in paras)
    return lines * pt * LINE_H_CM_PER_PT


def cap_h_cm(caption, width_cm, caption_pt) -> float:
    cpl = max(1, int(width_cm / (caption_pt * CHAR_W_CM_PER_PT)))
    lines = -(-int(eff_len(caption)) // cpl)
    return lines * caption_pt * LINE_H_CM_PER_PT + 0.15


def aspect_of(deps, fig_dir, fname) -> float:
    path = os.path.join(fig_dir, fname)
    with deps.Image.open(path) as img:
        width, height = img.size
    if width == 0:
        raise PosterError("overflow", f"figure {fname!r} has zero width")
    return height / width


def row_layout(deps, figs, fig_dir, inner_w, fig_w, caption_pt):
    metas, row_h, x = [], 0.0, 0.0
    for fig in figs:
        fh = fig_w * aspect_of(deps, fig_dir, fig["file"])
        ch = cap_h_cm(fig["caption"], fig_w, caption_pt)
        metas.append((fig, x, 0.0, fig_w))
        row_h = max(row_h, fh + ch)
        x += fig_w + FIG_GAP_CM
    return row_h, metas


def stacked_layout(deps, figs, fig_dir, inner_w, fig_w, caption_pt):
    metas, y = [], 0.0
    x = (inner_w - fig_w) / 2
    for fig in figs:
        fh = fig_w * aspect_of(deps, fig_dir, fig["file"])
        ch = cap_h_cm(fig["caption"], fig_w, caption_pt)
        metas.append((fig, x, y, fig_w))
        y += fh + ch + 0.3
    return y - 0.3 if figs else 0.0, metas


def place_figures(slide, figs, fig_dir, left_cm, top_cm, box_w_cm, budget_cm,
                  deps, body_pt, font_name=DEFAULT_FONT):
    """Place figures below the text. Tries a side-by-side row and a stacked
    full-width layout, keeping whichever fills the leftover space best
    without overflowing. Returns (height_used_cm, files) or (None, reason)."""
    n = len(figs)
    if n == 0:
        return 0.0, []
    caption_pt = body_pt
    inner_w = box_w_cm - 2 * SIDE_MARGIN_CM
    row_w = (inner_w - (n - 1) * FIG_GAP_CM) / n

    fitting = []
    modes = [(row_layout, row_w)]
    if n > 1:
        modes.append((stacked_layout, inner_w))
    for fn, base_w in modes:
        w = base_w
        for _ in range(40):
            h, metas = fn(deps, figs, fig_dir, inner_w, w, caption_pt)
            if h <= budget_cm:
                if w >= 8.0:
                    fitting.append((h, metas))
                break
            w *= 0.97
    if not fitting:
        return None, f"figures do not fit in {budget_cm:.1f}cm at readable size"
    height, metas = max(fitting, key=lambda c: c[0])

    used_w = max(x + w for _, x, _, w in metas)
    x0 = left_cm + (box_w_cm - used_w) / 2
    placed = []
    for fig, x_off, y_off, w in metas:
        fh = w * aspect_of(deps, fig_dir, fig["file"])
        slide.shapes.add_picture(
            os.path.join(fig_dir, fig["file"]),
            cm(deps, x0 + x_off), cm(deps, top_cm + y_off),
            width=cm(deps, w))
        cap = slide.shapes.add_textbox(
            cm(deps, x0 + x_off), cm(deps, top_cm + y_off + fh + 0.05),
            cm(deps, w), cm(deps, cap_h_cm(fig["caption"], w, caption_pt)))
        cap.text_frame.word_wrap = True
        p = cap.text_frame.paragraphs[0]
        p.line_spacing = 1.0
        p.alignment = deps.PP_ALIGN.CENTER
        run = p.add_run()
        run.text = fig["caption"]
        set_run_font(run, caption_pt, deps, font_name=font_name)
        placed.append(fig["file"])
    return height, placed


def fill_header(header, title, authors, title_ph, authors_ph):
    for para in header.text_frame.paragraphs:
        text = "".join(r.text for r in para.runs)
        if title_ph and title_ph in text:
            para.runs[0].text = title
            for run in para.runs[1:]:
                run.text = ""
        elif authors_ph and authors_ph in text:
            para.runs[0].text = authors
            for run in para.runs[1:]:
                run.text = ""


def build_parser() -> argparse.ArgumentParser:
    ap = JsonArgumentParser(
        description="Fill a poster PPTX form from a markdown content file.")
    ap.add_argument("--content", required=True)
    ap.add_argument("--figures", required=True)
    ap.add_argument("--form", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument(
        "--box-map",
        help="JSON object or path to JSON file mapping box keys to "
             "[content_shape, label_shape_or_null]. Overrides a box_map: "
             "header in the content file.",
    )
    ap.add_argument("--min-pt", type=float, default=24.0,
                    help="body/caption point size and overflow floor (default 24)")
    ap.add_argument("--header-shape", default=None)
    ap.add_argument("--title-placeholder", default=None)
    ap.add_argument("--authors-placeholder", default=None)
    ap.add_argument("--font", default=DEFAULT_FONT)
    return ap


def run_build(args, deps) -> dict:
    content = parse_content(args.content)
    if args.box_map:
        box_map = load_box_map_arg(args.box_map)
    elif content["box_map"]:
        box_map = content["box_map"]
    else:
        box_map = dict(DEFAULT_BOX_MAP)

    form_path = Path(args.form).resolve()
    out_path = Path(args.out).resolve()
    if form_path == out_path:
        raise PosterError("invalid_paths", "--out must not be the same file as --form")

    tmp_out = str(out_path) + ".tmp.pptx"
    shutil.copyfile(form_path, tmp_out)
    prs = None
    try:
        prs = deps.Presentation(tmp_out)
        if not prs.slides:
            raise PosterError("invalid_form", "form PPTX has no slides")
        slide = prs.slides[0]
        by_name = {sh.name: sh for sh in slide.shapes}
        body_pt = args.min_pt

        removed = 0
        for sh in list(slide.shapes):
            if sh.has_text_frame and any(
                    re.search(p, sh.text_frame.text) for p in GUIDE_PATTERNS):
                sh._element.getparent().remove(sh._element)
                removed += 1

        header_name = (
            args.header_shape
            or content["header_shape"]
            or DEFAULT_HEADER_SHAPE
        )
        header = by_name.get(header_name)
        if header is not None and header.has_text_frame:
            fill_header(
                header,
                content["title"],
                content["authors"],
                args.title_placeholder or content["title_placeholder"] or DEFAULT_TITLE_PLACEHOLDER,
                args.authors_placeholder or content["authors_placeholder"] or DEFAULT_AUTHORS_PLACEHOLDER,
            )

        failures = []
        box_stats = {}
        boxes = content["boxes"]
        for key, (box_name, label_name) in box_map.items():
            if key not in boxes:
                failures.append(f"{key}: missing from content file")
                continue
            box = by_name.get(box_name)
            if box is None or not box.has_text_frame:
                failures.append(f"{key}: content shape {box_name!r} not in form")
                continue
            payload = boxes[key]
            box_l = box.left / EMU_PER_CM
            box_t = box.top / EMU_PER_CM
            box_w = box.width / EMU_PER_CM
            box_h = box.height / EMU_PER_CM
            label_bottom = box_t
            if label_name and label_name in by_name:
                lab = by_name[label_name]
                label_bottom = (lab.top + lab.height) / EMU_PER_CM
            text_top = max(box_t, label_bottom) + 0.3
            usable_h = box_t + box_h - 0.3 - text_top
            text_w = box_w - 2 * SIDE_MARGIN_CM

            fill_text(box.text_frame, payload["paras"], body_pt, deps,
                      font_name=args.font)
            box.text_frame.margin_top = cm(deps, text_top - box_t)
            box.text_frame.margin_left = cm(deps, SIDE_MARGIN_CM)
            box.text_frame.margin_right = cm(deps, SIDE_MARGIN_CM)

            text_h = est_text_h_cm(payload["paras"], text_w, body_pt)
            budget = usable_h - text_h - (0.4 if payload["figs"] else 0)
            stats = {
                "text_cm": round(text_h, 2),
                "usable_cm": round(usable_h, 2),
                "figs": [],
            }
            if budget < 0:
                failures.append(
                    f"{key}: text {text_h:.1f}cm > box {usable_h:.1f}cm "
                    f"at {body_pt:g}pt (overflow)"
                )
                box_stats[key] = stats
                continue
            res, info = place_figures(
                slide, payload["figs"], args.figures,
                box_l, text_top + text_h + 0.4, box_w, budget,
                deps, body_pt, font_name=args.font)
            if res is None:
                failures.append(f"{key}: {info} (overflow)")
            else:
                stats["figs"] = info
            box_stats[key] = stats

        if failures:
            raise PosterError(
                "overflow",
                "poster overflow: " + "; ".join(failures),
                failures=failures,
            )

        prs.save(tmp_out)
        prs = None
        os.replace(tmp_out, out_path)
        tmp_out = None
        return {
            "ok": True,
            "out": str(out_path),
            "title": content["title"],
            "authors": content["authors"],
            "min_pt": body_pt,
            "guide_boxes_removed": removed,
            "boxes": box_stats,
        }
    finally:
        prs = None
        if tmp_out and os.path.exists(tmp_out):
            try:
                os.remove(tmp_out)
            except OSError:
                pass


def main(argv=None) -> int:
    _utf8_stdio()
    args = build_parser().parse_args(argv)
    deps = import_poster_deps()
    try:
        payload = run_build(args, deps)
    except PosterError as exc:
        error = {"code": exc.code, "message": exc.message}
        if exc.extra is not None:
            error["extra"] = exc.extra
        body = {"ok": False, "error": error}
        if exc.failures:
            body["failures"] = exc.failures
        code = 3 if exc.code == "poster_dependency_missing" else 1
        if exc.code in {"usage", "invalid_content", "invalid_box_map", "invalid_paths"}:
            code = 2
        _print_json(body)
        return code
    _print_json(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
