#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verification gates for a built poster PPTX.

Exits 1 on any failed gate. Prints exactly one JSON object on stdout.
Requires the optional extra: pip install ".[poster]"
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from types import SimpleNamespace


DEFAULT_FONT = "맑은 고딕"
DEFAULT_FORBIDDEN = (
    r"본문 내용은",
    r"맑은 고딕\s*\d+\s*pt",
)


class JsonArgumentParser(argparse.ArgumentParser):
    def error(self, message):
        _die({"ok": False, "error": {"code": "usage", "message": message}}, 2)


def _print_json(obj: dict) -> None:
    data = json.dumps(obj, ensure_ascii=False) + "\n"
    try:
        sys.stdout.write(data)
    except UnicodeEncodeError:
        sys.stdout.buffer.write(data.encode("utf-8"))
    sys.stdout.flush()


def _die(payload: dict, code: int) -> None:
    _print_json(payload)
    raise SystemExit(code)


def _utf8_stdio() -> None:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError, OSError):
            pass


def _missing_dep(exc: BaseException) -> None:
    _die(
        {
            "ok": False,
            "error": {
                "code": "poster_dependency_missing",
                "message": (
                    "python-pptx is required "
                    f"({type(exc).__name__}: {exc}). "
                    'Install with pip install ".[poster]".'
                ),
                "extra": "poster",
            },
        },
        3,
    )


def import_poster_deps():
    try:
        from pptx import Presentation
        from pptx.util import Pt
        from PIL import Image  # noqa: F401 — extra [poster] is pptx+Pillow
        try:
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            picture_type = MSO_SHAPE_TYPE.PICTURE
        except Exception:
            picture_type = 13
    except ImportError as exc:
        _missing_dep(exc)
    return SimpleNamespace(
        Presentation=Presentation, Pt=Pt, picture_type=picture_type)


def all_text(slide) -> str:
    chunks = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            chunks.append(sh.text_frame.text)
    return "\n".join(chunks)


def build_parser() -> argparse.ArgumentParser:
    ap = JsonArgumentParser(description="Verification gates for a built poster PPTX.")
    ap.add_argument("--out", required=True)
    ap.add_argument("--form", required=True)
    ap.add_argument(
        "--form-mtime", type=float, required=True,
        help="mtime of the form recorded before the build",
    )
    ap.add_argument(
        "--numbers",
        help="semicolon-separated key figures that must appear verbatim "
             "(semicolon, because figures like 4,182,797 contain commas)",
    )
    ap.add_argument(
        "--no-numbers", action="store_true",
        help="explicitly skip the key-number gates",
    )
    ap.add_argument("--min-pt", type=float, default=24.0)
    ap.add_argument("--max-pt", type=float, default=30.0)
    ap.add_argument(
        "--min-pictures", type=int, default=0,
        help="minimum picture shapes on the slide (default 0; pass a floor "
             "when the form is known to ship logos plus inserted figures)",
    )
    ap.add_argument(
        "--require-text", action="append", default=[],
        help="string that must appear in the poster (repeatable)",
    )
    ap.add_argument(
        "--forbidden", action="append", default=None,
        help="regex that must not appear; if omitted, form-guide defaults apply",
    )
    ap.add_argument(
        "--skip-shapes", action="append", default=[],
        help="shape names exempt from the body-font gate (repeatable)",
    )
    ap.add_argument("--font", default=DEFAULT_FONT)
    return ap


def run_verify(args, deps) -> tuple[dict, int]:
    nums = [n.strip() for n in (args.numbers or "").split(";") if n.strip()]
    if not nums and not args.no_numbers:
        _die({
            "ok": False,
            "error": {
                "code": "usage",
                "message": "--numbers is required (pass --no-numbers to skip "
                           "key-number gates on purpose)",
            },
        }, 2)

    forbidden = args.forbidden if args.forbidden is not None else list(DEFAULT_FORBIDDEN)
    skip_shapes = set(args.skip_shapes or [])
    min_pt = args.min_pt
    max_pt = max(args.max_pt, min_pt)
    results = []

    def gate(name, ok, detail=""):
        results.append({"name": name, "passed": bool(ok), "detail": detail})

    form_mtime = os.path.getmtime(args.form)
    gate("form mtime unchanged",
         abs(form_mtime - args.form_mtime) < 1e-6,
         str(form_mtime))

    prs = deps.Presentation(args.out)
    if not prs.slides:
        gate("poster has a slide", False, "no slides")
        payload = {
            "ok": False,
            "error": {"code": "verify_failed", "message": "poster has no slides"},
            "gates": results,
        }
        return payload, 1
    slide = prs.slides[0]
    text = all_text(slide)

    for pat in forbidden:
        gate(f"no /{pat}/", not re.search(pat, text))

    for required in args.require_text:
        gate(f"required text present: {required}", required in text)

    pics = [sh for sh in slide.shapes if sh.shape_type == deps.picture_type]
    gate(
        f"pictures >= {args.min_pictures}",
        len(pics) >= args.min_pictures,
        f"pictures={len(pics)}",
    )

    bad_fonts = []
    for sh in slide.shapes:
        if sh.name in skip_shapes:
            continue
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if not run.text.strip():
                    continue
                sz = run.font.size
                nm = run.font.name
                if sz is None:
                    continue
                if sz > deps.Pt(max_pt):
                    continue
                if sz < deps.Pt(min_pt):
                    bad_fonts.append((sh.name, run.text[:15], str(sz), nm))
                    continue
                if nm not in (None, args.font):
                    bad_fonts.append((sh.name, run.text[:15], str(sz), nm))
    gate(
        f"body fonts {args.font} {min_pt:g}-{max_pt:g}pt",
        not bad_fonts,
        str(bad_fonts[:5]),
    )

    for num in nums:
        gate(f"key number intact: {num}", num in text)

    ok = all(row["passed"] for row in results)
    if ok:
        return {"ok": True, "gates": results}, 0
    failed = [row["name"] for row in results if not row["passed"]]
    return {
        "ok": False,
        "error": {
            "code": "verify_failed",
            "message": "failed gates: " + "; ".join(failed),
        },
        "gates": results,
    }, 1


def main(argv=None) -> int:
    _utf8_stdio()
    args = build_parser().parse_args(argv)
    deps = import_poster_deps()
    payload, code = run_verify(args, deps)
    _print_json(payload)
    return code


if __name__ == "__main__":
    raise SystemExit(main())
